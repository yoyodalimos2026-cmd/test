from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path
from typing import Callable
from uuid import uuid4


@dataclass(frozen=True)
class SheetPreview:
    name: str
    rows: tuple[tuple[str, ...], ...]


@dataclass(frozen=True)
class WorkbookPreview:
    path: Path
    sheets: tuple[SheetPreview, ...]


def _working_copy_path(source: Path) -> Path:
    return source.with_name(f"{source.stem}_travail_{uuid4().hex}.xlsm")


def inspect_workbook(
    path: Path,
    max_rows: int = 30,
    max_columns: int = 16,
    progress_callback: Callable[[str], None] | None = None,
) -> WorkbookPreview:
    """Read a bounded preview without changing the selected workbook."""
    from openpyxl import load_workbook

    path = path.expanduser().resolve()
    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise ValueError("Formats pris en charge : .xlsx, .xlsm, .xltx et .xltm.")
    if progress_callback:
        progress_callback("Ouverture du classeur en lecture seule…")
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_vba=path.suffix.lower() in {".xlsm", ".xltm"},
    )
    try:
        previews = []
        total_sheets = len(workbook.worksheets)
        for index, sheet in enumerate(workbook.worksheets, start=1):
            if progress_callback:
                progress_callback(f"Lecture de l’aperçu : feuille {index}/{total_sheets} — {sheet.title}")
            rows = tuple(
                tuple("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=min(sheet.max_row or 1, max_rows),
                    min_col=1,
                    max_col=min(sheet.max_column or 1, max_columns),
                    values_only=True,
                )
            )
            previews.append(SheetPreview(sheet.title, rows))
        return WorkbookPreview(path, tuple(previews))
    finally:
        workbook.close()


def run_vba_macro(
    source: Path,
    module_path: Path,
    procedure: str,
    output_path: Path,
    progress_callback: Callable[[str], None] | None = None,
) -> Path:
    """Import a module set and run one macro in a separate Excel output copy."""
    if os.name != "nt":
        raise RuntimeError("L’exécution VBA nécessite Microsoft Excel sous Windows.")
    source = source.expanduser().resolve()
    module_path = module_path.expanduser().resolve()
    output = output_path.expanduser().resolve()
    module_folder = module_path.parent
    if source.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise ValueError("Excel peut uniquement traiter les classeurs .xlsx, .xlsm, .xltx ou .xltm ici.")
    if not procedure.strip():
        raise ValueError("Aucune procédure VBA sélectionnée.")
    if output.suffix.lower() != ".xlsm":
        raise ValueError("Le fichier de sortie doit avoir l’extension .xlsm pour conserver les macros.")
    if output == source:
        raise ValueError("Le fichier de sortie doit être différent du classeur original.")
    if not output.parent.is_dir():
        raise ValueError("Le dossier choisi pour le fichier de sortie n’existe pas.")
    if output.exists():
        raise FileExistsError(f"Le fichier de sortie existe déjà : {output}")
    modules = sorted(module_folder.glob("*.bas")) + sorted(module_folder.glob("*.frm"))
    if not modules:
        raise ValueError(f"Aucun module VBA trouvé dans {module_folder}.")
    if module_path not in {module.resolve() for module in modules}:
        raise ValueError(f"Le module sélectionné est absent de {module_folder}.")
    try:
        import pythoncom
        import win32com.client
    except ImportError as error:
        raise RuntimeError("Installez pywin32 sous Windows pour piloter Excel.") from error

    working_copy = _working_copy_path(source)
    excel = None
    workbook = None
    succeeded = False
    com_initialized = False

    def report(message: str) -> None:
        if progress_callback:
            progress_callback(message)

    try:
        pythoncom.CoInitialize()
        com_initialized = True
        report("Démarrage d’Excel…")
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AutomationSecurity = 3
        report("Ouverture du classeur complet…")
        workbook = excel.Workbooks.Open(str(source), UpdateLinks=0, ReadOnly=False)
        report("Création de la copie de travail…")
        workbook.SaveAs(str(working_copy), FileFormat=52)
        imported_modules = {}
        for index, module in enumerate(modules, start=1):
            report(f"Import du module VBA {index}/{len(modules)}…")
            imported_modules[module.resolve()] = workbook.VBProject.VBComponents.Import(str(module)).Name
        excel.AutomationSecurity = 1
        component_name = imported_modules[module_path]
        report(f"Exécution de la procédure {procedure}…")
        excel.Run(f"'{working_copy.name}'!{component_name}.{procedure}")
        excel.AutomationSecurity = 3
        excel.EnableEvents = False
        report("Enregistrement du classeur complet à la destination choisie…")
        workbook.SaveAs(str(output), FileFormat=52)
        workbook.Save()
        succeeded = True
        report("Classeur enregistré. Fermeture d’Excel…")
        return output
    except Exception as error:
        raise RuntimeError(
            "Excel n’a pas pu charger ou exécuter cette macro. Vérifiez que l’accès au modèle d’objet VBA "
            "est autorisé dans le Centre de gestion de la confidentialité, et que le classeur contient "
            f"les feuilles attendues. Détail : {error}"
        ) from error
    finally:
        try:
            try:
                if workbook is not None:
                    workbook.Close(SaveChanges=False)
            finally:
                if excel is not None:
                    excel.Quit()
        finally:
            try:
                if working_copy.exists():
                    working_copy.unlink(missing_ok=True)
                if output.exists() and not succeeded:
                    output.unlink(missing_ok=True)
            finally:
                if com_initialized:
                    pythoncom.CoUninitialize()
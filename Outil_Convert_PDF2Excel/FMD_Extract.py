import pdfplumber
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, Side
import re
import os
import platform
import sys
from tkinter import Tk, Label, Button, filedialog, StringVar, IntVar, PhotoImage, Canvas, messagebox
from tkinter.ttk import Progressbar
import threading

def clean_text(text):
    if re.search(r'\d', text):
        return text
    return re.sub(r'(.)\1+', r'\1', text)

def extract_table_from_pdf(pdf_path):
    try:
        pattern = r"\b\d{4}-\d{2}-\d{1}\b|\b\d{5}-\d{2}-\d{2}\b"
        with pdfplumber.open(pdf_path) as pdf:
            combined_df = pd.DataFrame()
            for page_number in range(len(pdf.pages)):
                page = pdf.pages[page_number]
                table = page.extract_table()
                if table:
                    df = pd.DataFrame(table[1:], columns=table[0])
                    df = df.applymap(lambda x: clean_text(str(x)) if isinstance(x, str) else x)
                    new_column_names = [
                        f"{old}-{new}" if old not in (None, "", " ") else new
                        for old, new in zip(df.columns, df.iloc[0])
                    ]
                    df.columns = new_column_names
                    df = df[1:]
                    if any(df.apply(lambda row: any(re.search(pattern, str(cell)) for cell in row), axis=1)):
                        combined_df = pd.concat([combined_df, df],ignore_index=True)
            return combined_df
    except Exception as e:
        print(f"Error extracting table from PDF: {e}")
        return None
def extract_table_from_pdf_TE(pdf_path):
    try :
        with pdfplumber.open(pdf_path) as pdf:
            combined_df = pd.DataFrame()

            for page_number in range(len(pdf.pages)):
                if page_number==0 :
                    page = pdf.pages[page_number]
                    table = page.extract_table()
                    df = pd.DataFrame(table)
                    level_row_index = df[df.eq('Level').any(axis=1)].index[0]

                    df.columns = df.iloc[level_row_index]

                    df = df.iloc[level_row_index + 1:]

                    df = df.reset_index(drop=True)
                    df = df.loc[:, df.columns.notnull()]
                else :
                    page = pdf.pages[page_number]
                    table = page.extract_table()
                    df = pd.DataFrame(table)



                    df = df.loc[:, df.columns != '']
                    #print (df)
                    df.columns=combined_df.columns
                combined_df = pd.concat([combined_df, df],axis=0 )

        return combined_df
    except Exception as e:
        print(f"Error extracting table from PDF: {e}")
        return None

def add_borders_to_cells(ws):
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin_border

def save_to_excel_with_borders(df, excel_path):
    df.to_excel(excel_path, index=False, engine='openpyxl')
    wb = load_workbook(excel_path)
    ws = wb.active
    add_borders_to_cells(ws)
    wb.save(excel_path)

def process_pdfs_in_directory(directory, progress_var, status_var):
    pdf_files = [f for f in os.listdir(directory) if f.endswith('.pdf')]
    total_files = len(pdf_files)

    for i, filename in enumerate(pdf_files):
        status_var.set(f"Extraction en cours : {filename}")
        pdf_path = os.path.join(directory, filename)
        excel_path = os.path.join(directory, filename.replace('.pdf', '.xlsx'))

        if os.path.basename(pdf_path).startswith('TE'):

            df = extract_table_from_pdf_TE(pdf_path)
        else:

            df = extract_table_from_pdf(pdf_path)
        if df is not None and not df.empty:
            save_to_excel_with_borders(df, excel_path)
            print(f'Extraction done for {filename}')
        else:
            print(f'No matching table found in {filename}')

        progress_var.set(int((i + 1) / total_files * 100))

    open_directory(directory)
    status_var.set("Extraction terminée")
    messagebox.showinfo("Terminé", "L'extraction est faite.")

def process_selected_pdfs(files, progress_var, status_var):
    total_files = len(files)

    for i, pdf_path in enumerate(files):
        status_var.set(f"Extraction en cours : {os.path.basename(pdf_path)}")
        excel_path = os.path.splitext(pdf_path)[0] + ".xlsx"
        if os.path.basename(pdf_path).startswith('TE'):

            df = extract_table_from_pdf_TE(pdf_path)
        else:

            df = extract_table_from_pdf(pdf_path)

        if df is not None and not df.empty:
            save_to_excel_with_borders(df, excel_path)
            print(f'Extraction done for {os.path.basename(pdf_path)}')
        else:
            print(f'No matching table found in {os.path.basename(pdf_path)}')

        progress_var.set(int((i + 1) / total_files * 100))

    directory = os.path.dirname(files[0])
    open_directory(directory)
    status_var.set("Extraction terminée")

def open_directory(directory):
    if platform.system() == "Windows":
        os.startfile(directory)
    elif platform.system() == "Darwin":
        os.system(f"open {directory}")
    else:
        os.system(f"xdg-open {directory}")

def resource_path(relative_path):

    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

def select_directory(progress_var, status_var, path_var):
    directory = filedialog.askdirectory()
    if directory:
        path_var.set(directory)
        threading.Thread(target=process_pdfs_in_directory, args=(directory, progress_var, status_var)).start()

def select_files(progress_var, status_var, path_var):
    files = filedialog.askopenfilenames(filetypes=[("PDF files", "*.pdf")])
    if files:
        path_var.set(", ".join(files))
        threading.Thread(target=process_selected_pdfs, args=(files, progress_var, status_var)).start()

def create_gui():
    root = Tk()
    root.title("PDF to Excel Converter")

    canvas = Canvas(root, width=400, height=300)
    canvas.pack()
    img_path = resource_path("background.png")
    img = PhotoImage(file=img_path)
    canvas.create_image(200, 200, anchor='center', image=img)

    path_var = StringVar()
    progress_var = IntVar(value=0)
    status_var = StringVar()

    Label(root, text="Convertir des FMD en Excel").pack(pady=10)
    Label(root, textvariable=path_var, wraplength=300).pack(pady=5)
    Button(root, text="sélectionner un Dossier", command=lambda: select_directory(progress_var, status_var, path_var)).pack(pady=10)
    Button(root, text="sélectionner un Fichier", command=lambda: select_files(progress_var, status_var, path_var)).pack(pady=10)
    Progressbar(root, length=300, variable=progress_var).pack(pady=10)
    Label(root, textvariable=status_var).pack(pady=5)

    root.mainloop()

if __name__ == '__main__':
    create_gui()

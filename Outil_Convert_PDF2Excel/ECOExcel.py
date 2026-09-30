import pdfplumber
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, Side,Font
import re
import os
import platform
import sys
from tkinter import Tk, Label, Button, filedialog, StringVar, IntVar, PhotoImage, Canvas, messagebox
from tkinter.ttk import Progressbar
import threading


def extract_table_from_pdf(pdf_path , excel_path):
    with pdfplumber.open(pdf_path) as pdf:
        # Extraire les tables de la seule page du PDF
        page = pdf.pages[0]
        tables = page.extract_tables()
        if not tables:
            print("Aucune table trouvée dans le fichier PDF.")
            return

        table = tables[0]

        indexed_lines = [row for row in table if any(cell.isdigit() for cell in row if cell)]

        df = pd.DataFrame(indexed_lines[1:],
                          columns=indexed_lines[0])

        df = df.dropna(axis=1, thresh=len(df) - 2)

        df = df.dropna(axis=0, thresh=len(df.columns) - 2)

        return df
def add_borders_to_cells(ws):
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'),
                         bottom=Side(style='thin'))
    plain_font = Font(bold=False)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin_border
            cell.font = plain_font


def save_to_excel_with_borders(df, excel_path):
    df.to_excel(excel_path, index=False, engine='openpyxl')
    wb = load_workbook(excel_path)
    ws = wb.active
    add_borders_to_cells(ws)
    wb.save(excel_path)


def process_pdfs_in_directory(directory, progress_var, status_var):
    pdf_files = [f for f in os.listdir(directory) if (f.endswith('.pdf') or f.endswith('.PDF') )]
    total_files = len(pdf_files)
    print(total_files)

    for i, filename in enumerate(pdf_files):
        status_var.set(f"Extraction en cours : {filename}")
        pdf_path = os.path.join(directory, filename)
        excel_path = os.path.splitext(pdf_path)[0] + ".xlsx"
        df=extract_table_from_pdf(pdf_path , excel_path)
        if df is not None and not df.empty:
            save_to_excel_with_borders(df, excel_path)
            print(f'Extraction done for {filename}')
        else:
            print(f'No matching table found in {filename}')

        progress_var.set(int((i + 1) / total_files * 100))

    open_directory(directory)
    status_var.set("Extraction terminée")
    messagebox.showinfo("Terminé", "L'extraction est faite.")
def process_selected_pdfs(files, progress_var, status_var):
    total_files = len(files)

    for i, pdf_path in enumerate(files):
        status_var.set(f"Extraction en cours : {os.path.basename(pdf_path)}")
        excel_path = os.path.splitext(pdf_path)[0] + ".xlsx"
        df=extract_table_from_pdf(pdf_path, excel_path)
        if df is not None and not df.empty:
            save_to_excel_with_borders(df, excel_path)
            print(f'Extraction done for {os.path.basename(pdf_path)}')
        else:
            print(f'No matching table found in {os.path.basename(pdf_path)}')

        progress_var.set(int((i + 1) / total_files * 100))

    directory = os.path.dirname(files[0])
    open_directory(directory)
    status_var.set("Extraction terminée")

def open_directory(directory):
    if platform.system() == "Windows":
        os.startfile(directory)
    elif platform.system() == "Darwin":
        os.system(f"open {directory}")
    else:
        os.system(f"xdg-open {directory}")

def resource_path(relative_path):
    """Obtenir le chemin des ressources incluses."""
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

def select_directory(progress_var, status_var, path_var):
    directory = filedialog.askdirectory()
    if directory:
        path_var.set(directory)
        threading.Thread(target=process_pdfs_in_directory, args=(directory, progress_var, status_var)).start()

def select_files(progress_var, status_var, path_var):
    files = filedialog.askopenfilenames(filetypes=[("PDF files", "*.pdf")])
    if files:
        path_var.set(", ".join(files))
        threading.Thread(target=process_selected_pdfs, args=(files, progress_var, status_var)).start()

def create_gui():
    root = Tk()
    root.title("PDF to Excel Converter")

    canvas = Canvas(root, width=400, height=300)
    canvas.pack()
    img_path = resource_path("background.png")
    img = PhotoImage(file=img_path)
    canvas.create_image(200, 200, anchor='center', image=img)

    path_var = StringVar()
    progress_var = IntVar(value=0)
    status_var = StringVar()

    Label(root, text="Sélectionnez Répertoire ou Fichiers avec des fichiers PDF").pack(pady=10)
    Label(root, textvariable=path_var, wraplength=300).pack(pady=5)
    Button(root, text="Sélectionnez le dossier", command=lambda: select_directory(progress_var, status_var, path_var)).pack(pady=10)
    Button(root, text="Sélectionnez le fichier", command=lambda: select_files(progress_var, status_var, path_var)).pack(pady=10)
    Progressbar(root, length=300, variable=progress_var).pack(pady=10)
    Label(root, textvariable=status_var).pack(pady=5)

    root.mainloop()

if __name__ == '__main__':
    create_gui()

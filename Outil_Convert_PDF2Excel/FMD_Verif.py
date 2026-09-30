import tkinter as tk
from tkinter import filedialog, messagebox, StringVar, IntVar, PhotoImage, Canvas
from tkinter import ttk
import pandas as pd
import os
import sys
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


def extract_info_from_excel_file_type_1(file_path, T1):
    try:
        df = pd.read_excel(file_path)
    except Exception as e:
        messagebox.showerror("Erreur", f"Erreur lors de la lecture du fichier {file_path}: {e}")
        return

    L = []
    somme = 0
    file_name = os.path.splitext(os.path.basename(file_path))[0]
    L.append(file_name)

    if 'CAS Number' not in df.columns:
        messagebox.showwarning("Avertissement", f"Le fichier {file_path} ne contient pas de colonne 'CAS Number'.")
        return

    for i in range(df.shape[0]):
        if pd.notna(df.at[i, 'CAS Number']):
            L.append(df.at[i, 'CAS Number'])
            L.append(df.iloc[i, -2])
            valeur = str(df.iloc[i, -2]).strip()

            if ',' in valeur:
                valeur = valeur.replace(',', '.')

            try:
                valeur_float = float(valeur)
            except ValueError:
                valeur_float = 0.0

            somme += valeur_float

    x = float(somme)
    if 98 < x < 99.9:
        msg = 'Warning : ' + str(x)
    elif x >= 99.9:
        msg = 'OK : ' + str(x)
    else:
        msg = 'NOT OK :  ' + str(x)

    L.insert(1, msg)

    T1.append(L)


def extract_info_from_excel_file_type_2(file_path, T2):
    try:
        df = pd.read_excel(file_path)
    except Exception as e:
        messagebox.showerror("Erreur", f"Erreur lors de la lecture du fichier {file_path}: {e}")
        return

    L = []
    i = 0
    n = df.shape[0]
    file_name = os.path.splitext(os.path.basename(file_path))[0]
    L.append(file_name)
    test_not_ok=True
    test_warning=True
    a=0
    m=0
    while i < n:
        s = 0
        j = i
        for k in range(j, n):
            if pd.notna(df.at[k, 'Substance']):
                L.append(df['CAS Number'][k])
                L.append(df[df.columns[-4]][k])
                valeur = str(df.iloc[k, -4]).strip()
                if ',' in valeur:
                    valeur = valeur.replace(',', '.')

                try:
                    valeur_float = float(valeur)
                except ValueError:
                    valeur_float = 0.0

                s += valeur_float
            else:
                if s != 0:
                    a=a+s
                    m=m+1
                    if 98 < s < 99.9:
                        msg = 'Warning'
                        test_warning=False
                    elif s >= 99.9:
                        msg = 'OK'
                    else:
                        msg = 'NOT OK'
                        test_not_ok=False
                    L.append(msg + ': ' + str(float(s)))
                    print('done')
                break
        i = k + 1

    if test_not_ok == False :
        out='NOT OK : '
    elif test_warning == False :
        out='Warning :  '
    else :
        out='OK : '
    L.insert(1,out +str(a/m) )
    T2.append(L)


def apply_styles(wb, sheet_name):
    fill_green = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")
    fill_orange = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
    fill_red = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

    sheet = wb[sheet_name]
    for row in sheet.iter_rows(min_row=2, max_col=sheet.max_column, max_row=sheet.max_row):
        for cell in row:
            if isinstance(cell.value, str):
                if 'NOT OK' in cell.value:
                    cell.fill = fill_red
                elif 'Warning' in cell.value:
                    cell.fill = fill_orange
                elif 'OK' in cell.value:
                    cell.fill = fill_green


def extract_table_TE(file_path, T):
    try:
        df = pd.read_excel(file_path)
        df = df.applymap(lambda x: x.replace('\n', '') if isinstance(x, str) else x)

    except Exception as e:
        #messagebox.showerror("Erreur", f"Erreur lors de la lecture du fichier {file_path}: {e}")
        return
    L = []
    i = 2
    n = df.shape[0]
    file_name = os.path.splitext(os.path.basename(file_path))[0]
    L.append(file_name)

    #df['Substance\nConcentration'] = df['Substance\nConcentration'].str.replace('\n', '', regex=False)
   # if 'Substance\nConcentratio\nn' in df.columns.tolist() :
       # df['Substance\nConcentratio\nn'] = df['Substance\nConcentratio\nn'].str.replace('\n', '', regex=False)
        #df['Substance\nConcentratio\nn'] = pd.to_numeric(df['Substance\nConcentratio\nn'], errors='coerce')
    df.rename(columns={'Substance\nConcentratio\nn': 'Substance\nConcentration'}, inplace=True)
    #print(df['Substance\nConcentration'])
    #df['Substance\nConcentration'] = df['Substance\nConcentration'].replace('\n', '', regex=False)
    df['Substance\nConcentration'] = pd.to_numeric(df['Substance\nConcentration'], errors='coerce')


    test_not_ok = True
    test_warning = True
    a = 0
    m = 0
    while i < n:
        s = 0
        j = i
        for k in range(j, n):
            if pd.notna(df.iloc[k, -6]):
                L.append(df['Substance\nCAS'][k])
                L.append(df.iloc[k, -5])

                s += df.iloc[k, -5]
                #print(s)
            else:
                if s != 0:
                    a = a + s
                    m = m + 1
                    if 98 < s < 99.9:
                        msg = 'Warning'
                        test_warning=False
                    elif s >= 99.9:
                        msg = 'OK'
                    else:
                        msg = 'NOT OK'
                        test_not_ok=False
                    L.append(msg + ': ' + str(float(s)))
                break
            if k == n - 1:
                if s != 0:
                    if 98 < s < 99.9:
                        msg = 'Warning'
                        test_warning = False
                    elif s >= 99.9:
                        msg = 'OK'
                    else:
                        msg = 'NOT OK'
                        test_not_ok=False
                    L.append(msg + ': ' + str(float(s)))
                    # print('*')
        i = k + 1

    if test_not_ok == False:
        out = 'NOT OK : '
    elif test_warning == False:
        out = 'Warning :  '
    else:
        out = 'OK : '
    L.insert(1, out + str(a / m))
    T.append(L)


def extract_onsemi(file_path, T):
    L = []
   # try:
    df = pd.read_excel(file_path)
   # except Exception as e:
       # messagebox.showerror("Erreur", f"Erreur lors de la lecture du fichier {file_path}: {e}")
        #return
    file_name = os.path.splitext(os.path.basename(file_path))[0]
    L.append(file_name)
    df = df.iloc[1:].reset_index(drop=True)
    df.columns = df.iloc[0]
    df = df[1:].reset_index(drop=True)
    df = df.iloc[:-2]
    df.rename(columns={'Homogenous Material': 'HomogenousMaterial'}, inplace=True)
    df.rename(columns={'Homogenous\nMaterial': 'HomogenousMaterial'}, inplace=True)
    print(df.columns)
    n = df.shape[0]
    i = 0
    test_not_ok=True
    test_warning=True
    a=0
    m=0
    while i < n:
        s = 0
        j = i
        for k in range(j, n):

            if pd.isna(df.at[k, 'CAS #']):
                L.append('n/a')
            else :
                L.append(df['CAS #'][k])
            L.append(df[df.columns[-1]][k])
            valeur = str(df.iloc[k, -1]).strip()
            if ',' in valeur:
                valeur = valeur.replace(',', '.')

            valeur_float = float(valeur)
            s += valeur_float

            if ((pd.isna(df.at[k, 'HomogenousMaterial'])) and (pd.notna(df.at[k + 1, 'HomogenousMaterial']))) or (
                    (pd.notna(df.at[k, 'HomogenousMaterial'])) and (k == n - 1)) or (
                    (pd.notna(df.at[k, 'HomogenousMaterial'])) and (pd.notna(df.at[k + 1, 'HomogenousMaterial']))):
                if (s != 0):
                    a = a + s
                    m = m + 1
                    if 99 < s < 99.9:
                        msg = 'Warning'
                        test_warning=False
                    elif s >= 99.9:
                        msg = 'OK'
                    else:
                        msg = 'NOT OK'
                        test_not_ok=False
                    L.append(msg + ': ' + str(float(s)))

                break
        i = k + 1

    if test_not_ok == False:
        out = 'NOT OK : '
    elif test_warning == False:
        out = 'Warning :  '
    else:
        out = 'OK : '
    L.insert(1, out + str(a / m))
    T.append(L)


def format_TI(files, progress_var, status_var, save_path):
    T1 = []
    T2 = []
    total_files = len(files)
    progress_var.set(0)

    for i, file_path in enumerate(files):
        try:
            extract_info_from_excel_file_type_1(file_path, T1)
            extract_info_from_excel_file_type_2(file_path, T2)
        except PermissionError:
            messagebox.showwarning("Permission Denied", f"Permission denied for file: {file_path}")
            continue
        progress_var.set((i + 1) / total_files * 100)
        status_var.set(f"Traitement du fichier {i + 1} sur {total_files}")

    if not T1 and not T2:
        messagebox.showwarning("Aucun fichier traité",
                               "Aucun fichier Excel trouvé ou les fichiers ne contiennent pas de données valides.")
        return

    df_final_1 = pd.DataFrame(T1)
    df_final_1.columns = ['FMD'] + ['Etat'] + [''] * (df_final_1.shape[1] - 2)

    df_final_2 = pd.DataFrame(T2)
    df_final_2.columns = ['FMD']+ ['Etat'] + [''] * (df_final_2.shape[1] - 2)

    # Ici, on utilise directement save_path sans demander à l'utilisateur de choisir
    try:
        with pd.ExcelWriter(save_path, engine='openpyxl') as writer:
            df_final_1.to_excel(writer, sheet_name='Component Level', index=False)
            df_final_2.to_excel(writer, sheet_name='Material Level', index=False)

        wb = load_workbook(save_path)

        apply_styles(wb, 'Component Level')
        apply_styles(wb, 'Material Level')

        wb.save(save_path)
        print(save_path)

        messagebox.showinfo("Succès", f"Le fichier Excel a été enregistré avec succès sous {save_path} !")
    except Exception as e:
        messagebox.showerror("Erreur", f"Une erreur s'est produite lors de l'enregistrement du fichier : {e}")

    status_var.set("Traitement terminé.")


def format_ONsemi(files, progress_var, status_var, save_path):
    T1 = []
    total_files = len(files)
    progress_var.set(0)

    for i, file_path in enumerate(files):

            extract_onsemi(file_path, T1)
            progress_var.set((i + 1) / total_files * 100)
            status_var.set(f"Traitement du fichier {i + 1} sur {total_files}")



    if not T1:
        messagebox.showwarning("Aucun fichier traité",
                               "Aucun fichier Excel trouvé ou les fichiers ne contiennent pas de données valides.")
        return

    df_final_1 = pd.DataFrame(T1)
    df_final_1.columns = ['FMD'] +['Etat'] + [''] * (df_final_1.shape[1] - 2)

    if not save_path:
        return

    try:
        with pd.ExcelWriter(save_path, engine='openpyxl') as writer:
            df_final_1.to_excel(writer, sheet_name='Total weight ', index=False)

        wb = load_workbook(save_path)

        apply_styles(wb, 'Total weight ')

        wb.save(save_path)

        messagebox.showinfo("Succès", f"Le fichier Excel a été enregistré avec succès sous {save_path} !")
    except Exception as e:
        messagebox.showerror("Erreur", f"Une erreur s'est produite lors de l'enregistrement du fichier : {e}")

    status_var.set("Traitement terminé.")


def format_TE(files, progress_var, status_var, save_path):
    T1 = []
    total_files = len(files)
    progress_var.set(0)

    for i, file_path in enumerate(files):
        try:
            extract_table_TE(file_path, T1)


        except PermissionError:
            messagebox.showwarning("Permission Denied", f"Permission denied for file: {file_path}")
            continue
        progress_var.set((i + 1) / total_files * 100)
        status_var.set(f"Traitement du fichier {i + 1} sur {total_files}")

    if not T1:
        messagebox.showwarning("Aucun fichier traité",
                               "Aucun fichier Excel trouvé ou les fichiers ne contiennent pas de données valides.")
        return

    df_final_1 = pd.DataFrame(T1)
    df_final_1.columns = ['FMD']+['Etat'] + [''] * (df_final_1.shape[1] - 2)

    if not save_path:
        return

    try:
        with pd.ExcelWriter(save_path, engine='openpyxl') as writer:
            df_final_1.to_excel(writer, sheet_name='Substance Concentration', index=False)

        wb = load_workbook(save_path)

        apply_styles(wb, 'Substance Concentration')

        wb.save(save_path)

        messagebox.showinfo("Succès", f"Le fichier Excel a été enregistré avec succès sous {save_path} !")
    except Exception as e:
        messagebox.showerror("Erreur", f"Une erreur s'est produite lors de l'enregistrement du fichier : {e}")

    status_var.set("Traitement terminé.")


def process_files(files, progress_var, status_var, format_type, folder_path):
    if format_type == 'FMD_TI':
        save_path = os.path.join(folder_path, 'result_TI.xlsx')
        format_TI(files, progress_var, status_var, save_path)
    elif format_type == 'FMD_TE':
        save_path = os.path.join(folder_path, 'result_TE.xlsx')
        format_TE(files, progress_var, status_var, save_path)
    elif format_type == 'FMD_ONsemi':
        save_path = os.path.join(folder_path, 'result_ONsemi.xlsx')
        format_ONsemi(files, progress_var, status_var, save_path)


def select_directory(progress_var, status_var, path_var, format_var):
    folder_path = filedialog.askdirectory(title="Choisissez le dossier contenant les fichiers Excel")
    if folder_path:
        path_var.set(folder_path)
        files = [os.path.join(folder_path, f) for f in os.listdir(folder_path)
                 if f.endswith('.xlsx') and not f.startswith('~$')]
        process_files(files, progress_var, status_var, format_var.get(), folder_path)



def resource_path(relative_path):
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)


def create_gui():
    root = tk.Tk()
    root.title("Extraction d'informations Excel")

    canvas = tk.Canvas(root, width=400, height=300)
    canvas.pack()
    img_path = resource_path("background.png")
    img = PhotoImage(file=img_path)
    canvas.create_image(200, 200, anchor='center', image=img)

    path_var = StringVar()
    progress_var = IntVar(value=0)
    status_var = StringVar()
    format_var = StringVar(value='FMD_TI')

    tk.Label(root, text="Sélectionnez un dossier contenant des fichiers Excel:").pack(pady=10)
    tk.Label(root, textvariable=path_var, wraplength=300).pack(pady=5)
    tk.Button(root, text="Sélectionner un dossier",
              command=lambda: select_directory(progress_var, status_var, path_var, format_var)).pack(pady=10)
    ttk.Progressbar(root, length=300, variable=progress_var).pack(pady=10)
    tk.Label(root, textvariable=status_var).pack(pady=5)

    tk.Label(root, text="Sélectionnez le format de fichier:").pack(pady=10)
    tk.Radiobutton(root, text="FMD_TI", variable=format_var, value='FMD_TI').pack(pady=5)
    tk.Radiobutton(root, text="FMD_ONsemi", variable=format_var, value='FMD_ONsemi').pack(pady=5)
    tk.Radiobutton(root, text="FMD_TE", variable=format_var, value='FMD_TE').pack(pady=5)

    root.mainloop()


if __name__ == "__main__":
    create_gui()
